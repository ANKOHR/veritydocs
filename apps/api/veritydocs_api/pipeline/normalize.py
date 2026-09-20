from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

import pymupdf as fitz
from openpyxl import load_workbook

from ..storage import ObjectStore
from .types import NormalizedDocument, NormalizedPage, TableBlock, TextBlock


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _pdf_page(
    page: fitz.Page, page_number: int, store: ObjectStore, key_prefix: str
) -> NormalizedPage:
    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
    image_key = f"{key_prefix}/rendered/page-{page_number:03d}.png"
    store.put_bytes(image_key, pixmap.tobytes("png"))
    scale_x = pixmap.width / page.rect.width if page.rect.width else 1.0
    scale_y = pixmap.height / page.rect.height if page.rect.height else 1.0
    blocks: list[TextBlock] = []
    for block in page.get_text("blocks"):
        x0, y0, x1, y1, text = block[:5]
        text = str(text).strip()
        if text:
            blocks.append(
                TextBlock(
                    text=text,
                    bbox=[x0 * scale_x, y0 * scale_y, x1 * scale_x, y1 * scale_y],
                )
            )
    return NormalizedPage(
        page_number=page_number,
        width=pixmap.width,
        height=pixmap.height,
        blocks=blocks,
        image_key=image_key,
        ocr_used=False,
        ocr_confidence=1.0 if blocks else 0.0,
    )


def normalize_pdf(
    document_id: str, filename: str, mime_type: str, data: bytes, store: ObjectStore
) -> NormalizedDocument:
    pdf = fitz.open(stream=data, filetype="pdf")
    if pdf.is_encrypted and not pdf.authenticate(""):
        raise ValueError("Encrypted PDFs are not supported without an explicit decryption boundary")
    if pdf.page_count == 0:
        raise ValueError("The PDF contains no pages")
    prefix = f"documents/{document_id}"
    pages = [_pdf_page(page, index + 1, store, prefix) for index, page in enumerate(pdf)]
    return NormalizedDocument(document_id, filename, mime_type, pages, {"format": "pdf"})


def _header_row(rows: list[list[str]]) -> int:
    keywords = {"unit", "unit id", "rent", "monthly rent", "status", "lease end", "amount", "date"}
    scores = [
        sum(
            1
            for cell in row
            if cell.lower() in keywords or any(k in cell.lower() for k in keywords)
        )
        for row in rows
    ]
    return (
        max(range(len(scores)), key=lambda index: scores[index])
        if scores and max(scores) > 0
        else 0
    )


def _sheet_page(title: str, rows: list[list[Any]], sheet_name: str) -> NormalizedPage:
    string_rows = [[_clean(cell) for cell in row] for row in rows]
    header_index = _header_row(string_rows)
    headers = string_rows[header_index] if string_rows else []
    data: list[list[Any]] = []
    for row in rows[header_index + 1 :]:
        values = [_clean(cell) for cell in row]
        if not any(values):
            continue
        if values[0].lower().startswith("total"):
            continue
        data.append(values[: len(headers)])
    blocks = [TextBlock(title, kind="heading", section=sheet_name)] if title else []
    blocks.append(
        TextBlock(f"Sheet: {sheet_name}; header row: {header_index + 1}", section=sheet_name)
    )
    return NormalizedPage(
        page_number=1,
        blocks=blocks,
        tables=[
            TableBlock(
                headers=headers, rows=data, source_range=f"{sheet_name}!row-{header_index + 1}"
            )
        ],
        ocr_confidence=1.0,
    )


def normalize_xlsx(
    document_id: str, filename: str, mime_type: str, data: bytes
) -> NormalizedDocument:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    pages: list[NormalizedPage] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        title = _clean(rows[0][0]) if rows and rows[0] else sheet.title
        pages.append(_sheet_page(title, rows, sheet.title))
    return NormalizedDocument(document_id, filename, mime_type, pages, {"format": "xlsx"})


def normalize_csv(
    document_id: str, filename: str, mime_type: str, data: bytes
) -> NormalizedDocument:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    return NormalizedDocument(
        document_id, filename, mime_type, [_sheet_page(filename, rows, "CSV")], {"format": "csv"}
    )


def normalize_image(
    document_id: str, filename: str, mime_type: str, data: bytes, store: ObjectStore
) -> NormalizedDocument:
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    key = f"documents/{document_id}/rendered/page-001.{Path(filename).suffix.lstrip('.') or 'png'}"
    store.put_bytes(key, data)
    page = NormalizedPage(
        page_number=1, width=image.width, height=image.height, image_key=key, ocr_confidence=0.0
    )
    return NormalizedDocument(document_id, filename, mime_type, [page], {"format": "image"})


def normalize_document(
    document_id: str, filename: str, mime_type: str, data: bytes, store: ObjectStore
) -> NormalizedDocument:
    extension = Path(filename).suffix.lower()
    if mime_type == "application/pdf" or extension == ".pdf":
        return normalize_pdf(document_id, filename, mime_type, data, store)
    if extension in {".xlsx", ".xlsm"}:
        return normalize_xlsx(document_id, filename, mime_type, data)
    if extension == ".csv" or mime_type == "text/csv":
        return normalize_csv(document_id, filename, mime_type, data)
    if mime_type.startswith("image/") or extension in {".png", ".jpg", ".jpeg", ".webp"}:
        return normalize_image(document_id, filename, mime_type, data, store)
    return NormalizedDocument(
        document_id,
        filename,
        mime_type,
        [NormalizedPage(1, blocks=[TextBlock("Unknown document format")])],
        {"format": "unknown"},
    )
