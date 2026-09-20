from __future__ import annotations

import io
from decimal import Decimal

import pymupdf as fitz
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import Settings, get_settings
from .models import CaseModel, DocumentModel
from .service import object_store, process_document
from .storage import sha256_bytes


def _pdf(title: str, lines: list[str]) -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((54, 58), title, fontsize=19, fontname="helv", color=(0.08, 0.16, 0.25))
    y = 96
    for line in lines:
        page.insert_text((54, y), line, fontsize=11, fontname="helv")
        y += 23
    result = document.tobytes()
    document.close()
    return result


def build_rent_roll() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "September Rent Roll"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = (
        "Acme Property Rent Roll | Reported monthly rent: £87,420 | Total units: 48 | Occupied units: 42"
    )
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Prepared for acquisition underwriting; headers intentionally begin below notes."
    sheet["A4"] = "Source system export"
    sheet["A13"] = "Unit"
    sheet["B13"] = "Status"
    sheet["C13"] = "Monthly Rent"
    sheet["D13"] = "Lease End"
    occupied = 42
    for index in range(48):
        row = 14 + index
        unit_id = "3B" if index == 2 else f"{index + 1:02d}{'A' if index % 2 == 0 else 'B'}"
        status = "Occupied" if index < occupied else "Vacant"
        rent = Decimal(2080) if status == "Occupied" else Decimal(0)
        if index == 0:
            rent += Decimal(60)
        sheet.cell(row, 1, unit_id)
        sheet.cell(row, 2, status)
        sheet.cell(row, 3, float(rent))
        # Deliberately leave one date blank so the review queue shows a real exception.
        if index != 2 and status == "Occupied":
            sheet.cell(row, 4, "2026-09-30")
    sheet.cell(62, 1, "TOTAL")
    sheet.cell(62, 3, 87420)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def demo_documents() -> list[tuple[str, str, bytes]]:
    return [
        (
            "rent_roll.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            build_rent_roll(),
        ),
        (
            "operating_statement.pdf",
            "application/pdf",
            _pdf(
                "Property Operating Statement",
                [
                    "Property: 32 New Street, London",
                    "Reporting period: 2024/25",
                    "Rental income: £87,420",
                    "Annual revenue: £1,049,040",
                    "Operating expenses: £336,200",
                    "Annual NOI: £712,840",
                ],
            ),
        ),
        (
            "loan_summary.pdf",
            "application/pdf",
            _pdf(
                "Loan Summary | Current Facilities",
                [
                    "Property: 32 New Street, London",
                    "Current Facilities",
                    "Loan balance: £4,180,000",
                    "Interest rate: 5.75%",
                    "Maturity: 2029-06-30",
                ],
            ),
        ),
        (
            "property_tax_scan.pdf",
            "application/pdf",
            _pdf(
                "Property Tax Notice",
                [
                    "Property: 32 New Street, London",
                    "Accounting period: 2023/24",
                    "Annual property tax: £91,400",
                    "Notice reference: PT-88421",
                ],
            ),
        ),
    ]


def contract_demo_documents() -> list[tuple[str, str, bytes]]:
    return [
        (
            "construction_contract.pdf",
            "application/pdf",
            _pdf(
                "Construction Services Agreement",
                [
                    "Project: Riverside Refurbishment",
                    "Contract reference: CON-2026-014",
                    "Contractor: Northbank Build Ltd",
                    "Contract value: £1,240,000",
                    "Commencement date: 2026-01-15",
                    "Completion date: 2026-12-15",
                    "Payment terms: 30 days",
                ],
            ),
        ),
        (
            "payment_application.pdf",
            "application/pdf",
            _pdf(
                "Application for Payment",
                [
                    "Project: Riverside Refurbishment",
                    "Contract reference: CON-2026-014",
                    "Application number: PA-009",
                    "Applicant: Northbank Build Ltd",
                    "Claimed amount: £1,270,000",
                    "Application period: October 2026",
                    "Payment due date: 2026-11-30",
                ],
            ),
        ),
    ]


def seed_acme_case(
    db: Session,
    case_id: str | None = None,
    settings: Settings | None = None,
) -> CaseModel:
    settings = settings or get_settings()
    store = object_store(settings)
    case = db.get(CaseModel, case_id) if case_id else None
    if case is None:
        case = CaseModel(
            name="Acme Property Acquisition",
            property_address="32 New Street, London",
            tenant_id="demo-tenant",
            status="processing",
        )
        db.add(case)
        db.commit()
        db.refresh(case)

    existing = {
        document.filename: document
        for document in db.scalars(select(DocumentModel).where(DocumentModel.case_id == case.id))
    }
    to_process: list[DocumentModel] = []
    for filename, mime_type, data in demo_documents():
        if filename in existing and existing[filename].status == "complete":
            continue
        digest = sha256_bytes(data)
        document = existing.get(filename)
        if document is None:
            key = f"documents/{case.id}/original/{filename}"
            store.put_bytes(key, data)
            document = DocumentModel(
                case_id=case.id,
                filename=filename,
                sha256=digest,
                mime_type=mime_type,
                size_bytes=len(data),
                storage_key=key,
                status="uploaded",
            )
            db.add(document)
            db.commit()
            db.refresh(document)
        to_process.append(document)

    for document in to_process:
        process_document(db, document.id, settings=settings, store=store)
    return db.get(CaseModel, case.id) or case


def seed_contract_case(
    db: Session,
    case_id: str | None = None,
    settings: Settings | None = None,
) -> CaseModel:
    settings = settings or get_settings()
    store = object_store(settings)
    case = db.get(CaseModel, case_id) if case_id else None
    if case is None:
        case = CaseModel(
            name="Riverside Contract Review",
            tenant_id="demo-tenant",
            status="processing",
        )
        db.add(case)
        db.commit()
        db.refresh(case)
    existing = {
        document.filename: document
        for document in db.scalars(select(DocumentModel).where(DocumentModel.case_id == case.id))
    }
    for filename, mime_type, data in contract_demo_documents():
        if filename in existing and existing[filename].status == "complete":
            continue
        document = existing.get(filename)
        if document is None:
            key = f"documents/{case.id}/original/{filename}"
            store.put_bytes(key, data)
            document = DocumentModel(
                case_id=case.id,
                filename=filename,
                sha256=sha256_bytes(data),
                mime_type=mime_type,
                size_bytes=len(data),
                storage_key=key,
                status="uploaded",
            )
            db.add(document)
            db.commit()
            db.refresh(document)
        process_document(db, document.id, settings=settings, store=store)
    return db.get(CaseModel, case.id) or case
